#!/usr/bin/env python3
"""
generate_pptx.py — Create a PowerPoint presentation from a JSON spec.

Usage:
  echo '<json>' | python3 generate_pptx.py
  python3 generate_pptx.py '<json_string>'

JSON schema:
  {
    "title": "Presentation Title",
    "subtitle": "Optional subtitle",
    "slides": [
      {
        "title": "Slide Title",
        "bullets": ["Point 1", "Point 2", "Point 3"],
        "speaker_notes": "Optional presenter notes"
      }
    ]
  }

Output: /workspace/agent/output/slides.pptx
"""
import json
import os
import sys

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
except ImportError:
    print("Error: python3-pptx not installed. Run: apt-get install -y python3-pptx", file=sys.stderr)
    sys.exit(1)

OUTPUT_DIR = '/workspace/agent/output'
OUTPUT_PATH = os.path.join(OUTPUT_DIR, 'slides.pptx')

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
LOGO_PATH = os.path.join(SCRIPT_DIR, 'assets', 'synapxe-logo.png')

# Synapxe brand palette — sampled from the company logo.
PURPLE  = RGBColor(0x5B, 0x0D, 0xF6)  # primary violet
MAGENTA = RGBColor(0xBB, 0x03, 0xF9)  # gradient / accent magenta
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
SLATE   = RGBColor(0x1A, 0x1A, 0x2E)  # body text on white
MUTED   = RGBColor(0xE3, 0xD6, 0xFB)  # subtitle text on purple


def no_line(shape):
    shape.line.fill.background()


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)

    # Full-bleed violet-to-magenta gradient background (Synapxe brand)
    bg = slide.background
    bg.fill.gradient()
    stops = bg.fill.gradient_stops
    stops[0].position = 0.0
    stops[0].color.rgb = PURPLE
    stops[1].position = 1.0
    stops[1].color.rgb = MAGENTA
    bg.fill.gradient_angle = 45.0

    # White card behind the logo so it stays legible over the gradient
    card_left, card_top = Inches(0.5), Inches(1.3)
    card_width, card_height = Inches(3.3), Inches(1.5)
    card = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        card_left, card_top, card_width, card_height,
    )
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    no_line(card)
    card.adjustments[0] = 0.08

    # Logo (transparent-background PNG, sized to fit inside the card)
    if os.path.exists(LOGO_PATH):
        logo_height = Inches(1.0)
        logo_width = Inches(2.8)
        slide.shapes.add_picture(
            LOGO_PATH,
            card_left + Inches(0.25), card_top + Inches(0.25),
            width=logo_width, height=logo_height,
        )

    # Title
    tx = slide.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(12.33), Inches(1.6))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(48)
    run.font.bold = True
    run.font.color.rgb = WHITE

    # Subtitle
    if subtitle:
        tx2 = slide.shapes.add_textbox(Inches(0.5), Inches(4.9), Inches(12.33), Inches(0.9))
        tf2 = tx2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.alignment = PP_ALIGN.LEFT
        run2 = p2.runs[0]
        run2.font.size = Pt(22)
        run2.font.color.rgb = MUTED


def add_content_slide(prs: Presentation, slide_title: str, bullets: list, notes: str) -> None:
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)

    # White background
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE

    # Purple header bar
    header = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0), Inches(0), prs.slide_width, Inches(1.25),
    )
    header.fill.solid()
    header.fill.fore_color.rgb = PURPLE
    no_line(header)

    # Slide title inside header
    tx = slide.shapes.add_textbox(Inches(0.35), Inches(0.18), Inches(12.63), Inches(0.88))
    tf = tx.text_frame
    p = tf.paragraphs[0]
    p.text = slide_title
    run = p.runs[0]
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = WHITE

    # Accent left strip for content area
    strip = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0), Inches(1.25), Inches(0.08), prs.slide_height - Inches(1.25),
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = MAGENTA
    no_line(strip)

    # Bullet content
    if bullets:
        content = slide.shapes.add_textbox(
            Inches(0.45), Inches(1.45), Inches(12.43), Inches(5.8),
        )
        tf = content.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = bullet
            p.space_before = Pt(10)
            run = p.runs[0]
            run.font.size = Pt(20)
            run.font.color.rgb = SLATE

    # Speaker notes
    if notes:
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = notes


def main():
    if len(sys.argv) > 1:
        raw = sys.argv[1]
    else:
        raw = sys.stdin.read().strip()

    if not raw:
        print("Error: no JSON input provided.", file=sys.stderr)
        sys.exit(1)

    try:
        spec = json.loads(raw)
    except json.JSONDecodeError as e:
        print(f"Error: invalid JSON — {e}", file=sys.stderr)
        sys.exit(1)

    title = spec.get('title', 'Presentation')
    subtitle = spec.get('subtitle', '')
    slides_data = spec.get('slides', [])

    if not slides_data:
        print("Error: 'slides' array is empty or missing.", file=sys.stderr)
        sys.exit(1)

    prs = Presentation()
    prs.slide_width  = Inches(13.33)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    add_title_slide(prs, title, subtitle)

    for s in slides_data:
        add_content_slide(
            prs,
            slide_title=s.get('title', 'Untitled'),
            bullets=s.get('bullets', []),
            notes=s.get('speaker_notes', ''),
        )

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    prs.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH} ({len(slides_data)} content slides + 1 title slide)")


if __name__ == '__main__':
    main()
